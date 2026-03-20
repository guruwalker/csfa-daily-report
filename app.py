"""
CSFA Report Web Interface
Flask application for generating and previewing reports with date selection.

Updated: Added /api/suspensions and /api/new-starters endpoints.
"""

from flask import Flask, render_template, request, send_file, jsonify, flash, redirect, url_for
from datetime import datetime, timedelta
import os
import logging
from pathlib import Path
import traceback

from generate_detailed_report import generate_detailed_report, ReportConfig
from main import Config as MainConfig, fetch_timesheet_data, fetch_orders_data, format_date_for_api, format_date_range
from holiday_checker import is_holiday

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET_KEY", "csfa-report-secret-key-change-in-production")

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

REPORTS_FOLDER = Path("generated_reports")
REPORTS_FOLDER.mkdir(exist_ok=True)


# ============================================================================
# HELPER FUNCTIONS
# ============================================================================

def validate_date(date_str: str) -> tuple[bool, str, datetime]:
    """Validate date string and return parsed datetime."""
    if not date_str:
        return False, "Date is required", None

    try:
        report_date = datetime.strptime(date_str, "%Y-%m-%d")

        if report_date.date() > datetime.now().date():
            return False, "Cannot generate reports for future dates", None

        if report_date.weekday() >= 5:
            return False, "Cannot generate reports for weekends (Saturday/Sunday)", None

        return True, "", report_date

    except ValueError:
        return False, "Invalid date format. Use YYYY-MM-DD", None


def get_report_summary(report_file: str) -> dict:
    """Extract summary information from generated report."""
    import pandas as pd
    from openpyxl import load_workbook

    try:
        if not os.path.exists(report_file):
            return {"error": "Report file not found"}

        df_summary = pd.read_excel(report_file, sheet_name="Day Summary")

        total_visited = int(df_summary["CUSTOMERS VISITED"].sum())
        active_reps = int(df_summary[df_summary["APP USAGE"] == "Used App"].shape[0])
        total_reps = len(df_summary)

        wb = load_workbook(report_file, read_only=True)
        sheets = wb.sheetnames
        wb.close()

        return {
            "total_customers_visited": total_visited,
            "active_salespeople": f"{active_reps}/{total_reps}",
            "sheets": sheets,
            "file_size": f"{os.path.getsize(report_file) / 1024:.1f} KB"
        }

    except Exception as e:
        logger.error(f"Error reading report: {e}")
        return {"error": str(e)}


# ============================================================================
# ROUTES
# ============================================================================

@app.route('/')
def index():
    """Home page with date selection form."""
    default_date = datetime.now().strftime("%Y-%m-%d")

    today = datetime.now()
    recent_dates = []

    for i in range(30):
        date = today - timedelta(days=i)
        if date.weekday() < 5:
            recent_dates.append({
                "date": date.strftime("%Y-%m-%d"),
                "display": date.strftime("%A, %B %d, %Y")
            })

    return render_template('index.html',
                          default_date=default_date,
                          recent_dates=recent_dates)


@app.route('/generate', methods=['POST'])
def generate_report():
    """Generate report for selected date."""
    try:
        report_date_str = request.form.get('report_date')

        is_valid, error_msg, report_date = validate_date(report_date_str)
        if not is_valid:
            flash(error_msg, 'error')
            return redirect(url_for('index'))

        logger.info(f"Generating report for {report_date_str}")

        if is_holiday(report_date):
            flash(f"{report_date_str} is a public holiday. Generating holiday report...", 'info')

        MainConfig.validate()

        order_date = format_date_for_api(report_date)
        order_date_range = format_date_range(report_date)

        logger.info("Fetching data from API...")
        flash("Fetching data from API...", 'info')

        orders_data = fetch_orders_data(order_date, order_date_range)
        visits_data = fetch_timesheet_data(order_date_range)

        logger.info("Generating Excel report...")
        flash("Generating Excel report...", 'info')

        report_config = ReportConfig.from_env()

        report_filename = f"CSFA_Report_{report_date_str}.xlsx"
        report_path = REPORTS_FOLDER / report_filename
        report_config.output_file = str(report_path)

        generate_detailed_report(visits_data, orders_data, report_date, report_config)

        flash("Report generated successfully!", 'success')

        summary = get_report_summary(str(report_path))

        return render_template('result.html',
                             report_date=report_date_str,
                             report_file=report_filename,
                             summary=summary,
                             email_sent=False)

    except Exception as e:
        logger.error(f"Error generating report: {e}")
        logger.error(traceback.format_exc())
        flash(f"Error generating report: {str(e)}", 'error')
        return redirect(url_for('index'))


@app.route('/download/<filename>')
def download_report(filename):
    """Download generated report."""
    try:
        report_path = REPORTS_FOLDER / filename

        if not report_path.exists():
            flash("Report file not found", 'error')
            return redirect(url_for('index'))

        return send_file(
            report_path,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    except Exception as e:
        logger.error(f"Error downloading report: {e}")
        flash(f"Error downloading report: {str(e)}", 'error')
        return redirect(url_for('index'))


@app.route('/preview/<filename>')
def preview_report(filename):
    """Preview report summary without downloading."""
    try:
        report_path = REPORTS_FOLDER / filename

        if not report_path.exists():
            return jsonify({"error": "Report file not found"}), 404

        summary = get_report_summary(str(report_path))

        import pandas as pd
        df = pd.read_excel(report_path, sheet_name="Day Summary")
        preview_html = df.head(10).to_html(classes='table table-striped', index=False)

        return jsonify({
            "summary": summary,
            "preview": preview_html
        })

    except Exception as e:
        logger.error(f"Error previewing report: {e}")
        return jsonify({"error": str(e)}), 500


@app.route('/send-email', methods=['POST'])
def send_email_route():
    """Send email for a generated report (test or live thread)."""
    try:
        data = request.get_json()
        filename = data.get('filename')
        mode = data.get('mode', 'test')
        report_date_str = data.get('report_date')

        if not filename:
            return jsonify({"success": False, "error": "Filename required"}), 400

        report_path = REPORTS_FOLDER / filename

        if not report_path.exists():
            return jsonify({"success": False, "error": "Report file not found"}), 404

        logger.info(f"Sending email in {mode} mode for {filename}")

        ENV_KEYS = ("EMAIL_TO", "EMAIL_CC", "EMAIL_THREAD_ID")
        original_env = {k: os.environ.get(k) for k in ENV_KEYS}

        def _set_or_remove(key: str, value):
            if value:
                os.environ[key] = value
            else:
                os.environ.pop(key, None)

        try:
            if mode == 'test':
                to = os.getenv('EMAIL_TO_TEST', os.getenv('SENDER_EMAIL', ''))
                cc = os.getenv('EMAIL_CC_TEST', '')
                thread_id = os.getenv('EMAIL_THREAD_ID_TEST', '')

                if not to:
                    return jsonify({"success": False, "error": "EMAIL_TO_TEST not configured in .env"}), 400

                _set_or_remove("EMAIL_TO", to)
                _set_or_remove("EMAIL_CC", cc)
                _set_or_remove("EMAIL_THREAD_ID", thread_id)
                logger.info(f"📧 Test mode → TO: {to}  THREAD: {thread_id or '(new thread)'}")

            else:
                to = os.getenv('EMAIL_TO_LIVE', '')
                cc = os.getenv('EMAIL_CC_LIVE', '')
                thread_id = os.getenv('EMAIL_THREAD_ID_LIVE', '')

                if not to:
                    return jsonify({"success": False, "error": "EMAIL_TO_LIVE not configured in .env"}), 400

                _set_or_remove("EMAIL_TO", to)
                _set_or_remove("EMAIL_CC", cc)
                _set_or_remove("EMAIL_THREAD_ID", thread_id)
                logger.info(f"📧 Live mode → TO: {to}  THREAD: {thread_id or '(new thread)'}")

            from send_report import send_report
            success = send_report(
                excel_file=str(report_path),
                summary_sheet="Day Summary",
                date_str=report_date_str
            )

        finally:
            for key, original_value in original_env.items():
                if original_value is None:
                    os.environ.pop(key, None)
                else:
                    os.environ[key] = original_value

        if success:
            return jsonify({
                "success": True,
                "mode": mode,
                "message": f"Email sent successfully in {mode} mode"
            })
        else:
            return jsonify({
                "success": False,
                "error": "Email sending failed - check server logs"
            }), 500

    except Exception as e:
        logger.error(f"Error sending email: {e}")
        logger.error(traceback.format_exc())
        return jsonify({"success": False, "error": str(e)}), 500


# ============================================================================
# HOLIDAYS & LEAVE ENDPOINTS (unchanged)
# ============================================================================

@app.route('/api/holidays-and-leave')
def get_holidays_and_leave():
    """Return current holidays and leave days for the UI."""
    import json
    from pathlib import Path

    holidays_file = Path(os.getenv("HOLIDAYS_FILE", "holidays.json"))
    leave_file = Path(os.getenv("LEAVE_FILE", "leave_days.json"))

    holidays = []
    if holidays_file.exists():
        try:
            holidays = json.loads(holidays_file.read_text()).get("holidays", [])
        except Exception:
            pass

    leave_days = {}
    if leave_file.exists():
        try:
            leave_days = json.loads(leave_file.read_text()).get("leave_days", {})
        except Exception:
            pass

    from config import get_all_salespeople
    return jsonify({
        "holidays": sorted(holidays),
        "leave_days": leave_days,
        "salespeople": get_all_salespeople()
    })


@app.route('/api/holidays-and-leave', methods=['POST'])
def update_holidays_and_leave():
    """Update holidays.json and/or leave_days.json from the UI."""
    import json
    from pathlib import Path

    data = request.get_json()
    action = data.get("action")
    date_str = data.get("date")
    person = data.get("person", "")

    if not action or not date_str:
        return jsonify({"success": False, "error": "action and date are required"}), 400

    try:
        datetime.strptime(date_str, "%Y-%m-%d")
    except ValueError:
        return jsonify({"success": False, "error": "Invalid date format"}), 400

    holidays_file = Path(os.getenv("HOLIDAYS_FILE", "holidays.json"))
    leave_file = Path(os.getenv("LEAVE_FILE", "leave_days.json"))

    try:
        if action == "add_holiday":
            obj = json.loads(holidays_file.read_text()) if holidays_file.exists() else {"holidays": []}
            if date_str not in obj["holidays"]:
                obj["holidays"].append(date_str)
                obj["holidays"].sort()
            holidays_file.write_text(json.dumps(obj, indent=2))
            return jsonify({"success": True})

        elif action == "remove_holiday":
            if not holidays_file.exists():
                return jsonify({"success": True})
            obj = json.loads(holidays_file.read_text())
            obj["holidays"] = [d for d in obj.get("holidays", []) if d != date_str]
            holidays_file.write_text(json.dumps(obj, indent=2))
            return jsonify({"success": True})

        elif action == "add_leave":
            if not person:
                return jsonify({"success": False, "error": "person is required for leave actions"}), 400
            obj = json.loads(leave_file.read_text()) if leave_file.exists() else {"leave_days": {}}
            if person not in obj["leave_days"]:
                obj["leave_days"][person] = []
            if date_str not in obj["leave_days"][person]:
                obj["leave_days"][person].append(date_str)
                obj["leave_days"][person].sort()
            leave_file.write_text(json.dumps(obj, indent=2))
            return jsonify({"success": True})

        elif action == "remove_leave":
            if not person:
                return jsonify({"success": False, "error": "person is required for leave actions"}), 400
            if not leave_file.exists():
                return jsonify({"success": True})
            obj = json.loads(leave_file.read_text())
            if person in obj.get("leave_days", {}):
                obj["leave_days"][person] = [d for d in obj["leave_days"][person] if d != date_str]
            leave_file.write_text(json.dumps(obj, indent=2))
            return jsonify({"success": True})

        else:
            return jsonify({"success": False, "error": f"Unknown action: {action}"}), 400

    except Exception as e:
        logger.error(f"Error updating holidays/leave: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ============================================================================
# SUSPENSIONS ENDPOINT
# ============================================================================

@app.route('/api/suspensions', methods=['GET'])
def get_suspensions():
    """Return current suspended salespeople."""
    import json
    from pathlib import Path
    from config import get_all_salespeople

    suspended_file = Path(os.getenv("SUSPENDED_FILE", "suspended_salespeople.json"))
    suspended = {}
    if suspended_file.exists():
        try:
            suspended = json.loads(suspended_file.read_text()).get("suspended", {})
        except Exception:
            pass

    return jsonify({
        "suspended": suspended,
        "salespeople": get_all_salespeople()
    })


@app.route('/api/suspensions', methods=['POST'])
def update_suspensions():
    """
    Add or remove a suspension.

    Expected JSON:
      { "action": "suspend",   "person": "JOHN DOE", "date": "2026-03-20" }
      { "action": "reinstate", "person": "JOHN DOE" }
    """
    import json
    from pathlib import Path

    data = request.get_json()
    action = data.get("action")   # "suspend" | "reinstate"
    person = data.get("person", "").strip()
    date_str = data.get("date", "")

    if not action or not person:
        return jsonify({"success": False, "error": "action and person are required"}), 400

    if action == "suspend" and not date_str:
        return jsonify({"success": False, "error": "date is required for suspend action"}), 400

    if date_str:
        try:
            datetime.strptime(date_str, "%Y-%m-%d")
        except ValueError:
            return jsonify({"success": False, "error": "Invalid date format (use YYYY-MM-DD)"}), 400

    suspended_file = Path(os.getenv("SUSPENDED_FILE", "suspended_salespeople.json"))

    try:
        obj = json.loads(suspended_file.read_text()) if suspended_file.exists() else {"suspended": {}}

        if action == "suspend":
            obj["suspended"][person] = date_str
            logger.info(f"Suspended {person} from {date_str}")

        elif action == "reinstate":
            obj["suspended"].pop(person, None)
            logger.info(f"Reinstated {person}")

        else:
            return jsonify({"success": False, "error": f"Unknown action: {action}"}), 400

        suspended_file.write_text(json.dumps(obj, indent=2))
        return jsonify({"success": True})

    except Exception as e:
        logger.error(f"Error updating suspensions: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ============================================================================
# NEW STARTERS ENDPOINT
# ============================================================================

@app.route('/api/new-starters', methods=['GET'])
def get_new_starters():
    """Return current new starter records."""
    import json
    from pathlib import Path
    from config import get_all_salespeople

    starters_file = Path(os.getenv("NEW_STARTERS_FILE", "new_starters.json"))
    new_starters = {}
    if starters_file.exists():
        try:
            new_starters = json.loads(starters_file.read_text()).get("new_starters", {})
        except Exception:
            pass

    return jsonify({
        "new_starters": new_starters,
        "salespeople": get_all_salespeople()
    })


@app.route('/api/new-starters', methods=['POST'])
def update_new_starters():
    """
    Add or remove a new starter record.

    Expected JSON:
      { "action": "add",    "person": "JOHN DOE", "date": "2026-03-20" }
      { "action": "remove", "person": "JOHN DOE" }
    """
    import json
    from pathlib import Path

    data = request.get_json()
    action = data.get("action")   # "add" | "remove"
    person = data.get("person", "").strip()
    date_str = data.get("date", "")

    if not action or not person:
        return jsonify({"success": False, "error": "action and person are required"}), 400

    if action == "add" and not date_str:
        return jsonify({"success": False, "error": "date is required for add action"}), 400

    if date_str:
        try:
            datetime.strptime(date_str, "%Y-%m-%d")
        except ValueError:
            return jsonify({"success": False, "error": "Invalid date format (use YYYY-MM-DD)"}), 400

    starters_file = Path(os.getenv("NEW_STARTERS_FILE", "new_starters.json"))

    try:
        obj = json.loads(starters_file.read_text()) if starters_file.exists() else {"new_starters": {}}

        if action == "add":
            obj["new_starters"][person] = date_str
            logger.info(f"Registered new starter {person} from {date_str}")

        elif action == "remove":
            obj["new_starters"].pop(person, None)
            logger.info(f"Removed new starter record for {person}")

        else:
            return jsonify({"success": False, "error": f"Unknown action: {action}"}), 400

        starters_file.write_text(json.dumps(obj, indent=2))
        return jsonify({"success": True})

    except Exception as e:
        logger.error(f"Error updating new starters: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ============================================================================
# HEALTH & ERROR HANDLERS
# ============================================================================

@app.route('/health')
def health():
    """Health check endpoint."""
    return jsonify({
        "status": "healthy",
        "timestamp": datetime.now().isoformat(),
        "reports_folder": str(REPORTS_FOLDER),
        "reports_count": len(list(REPORTS_FOLDER.glob("*.xlsx")))
    })


@app.errorhandler(404)
def not_found(e):
    return render_template('404.html'), 404


@app.errorhandler(500)
def server_error(e):
    return render_template('500.html'), 500


# ============================================================================
# MAIN
# ============================================================================

if __name__ == '__main__':
    try:
        MainConfig.validate()
        logger.info("✅ Configuration validated")
    except Exception as e:
        logger.error(f"❌ Configuration error: {e}")
        logger.warning("App will start but report generation may fail")

    port = int(os.getenv('PORT', 5000))
    debug = os.getenv('FLASK_DEBUG', 'False').lower() == 'true'

    logger.info(f"🚀 Starting CSFA Report Web Interface on port {port}")
    logger.info(f"📁 Reports will be saved to: {REPORTS_FOLDER.absolute()}")

    app.run(host='0.0.0.0', port=port, debug=debug)
